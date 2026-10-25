"""Build fifteen Persian, service-specific Word manuals from reviewed prose and real contracts.

Run with the bundled document runtime, after export_guide_contracts.py for each service.
Rendering is a separate required operation, never inferred from successful DOCX creation.
"""
from datetime import datetime, timezone
import json
from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs/word-source'
QA = ROOT / 'artifacts/guide-qa'
OUTPUT = ROOT / 'deliverables/word-guides'
FONT = Path('C:/Windows/Fonts')
SERVICES = {
    'personnel': ('سامانه پرسنلی', '01-personnel', 3001, 8001),
    'inventory': ('سامانه انبار', '02-inventory', 3002, 8002),
    'commerce': ('سامانه بازرگانی', '03-commerce', 3003, 8003),
}
KINDS = {
    'architecture': ('01', 'راهنمای معماری'),
    'database': ('02', 'راهنمای طراحی و دیاگرام پایگاه داده'),
    'backend': ('03', 'راهنمای ساختار بک اند'),
    'frontend': ('04', 'راهنمای ساختار فرانت اند'),
    'user_api': ('05', 'راهنمای کاربری و توضیح API ها'),
}
COMMON_NOTES = {
    'users': 'حساب‌های مستقل همین سامانه شامل نقش و هش رمز؛ password_hash هرگز در پاسخ کاربران ارائه نمی‌شود.',
    'auth_sessions': 'نشست‌های هشت ساعته همین سامانه؛ تنها هش توکن نگهداری می‌شود و خروج یا تغییر رمز آن را باطل می‌کند.',
    'audit_logs': 'سابقه عملیات با نام کاربر در زمان ثبت؛ user_id و entity_id مرجع متنی‌اند و کلید خارجی دیتابیس نیستند.',
}
FIELD_NOTES = {
    'id': 'شناسه UUID رکورد', 'created_at': 'زمان ایجاد به UTC', 'updated_at': 'زمان آخرین تغییر به UTC',
    'username': 'نام کاربری یکتا', 'full_name': 'نام نمایشی حساب', 'password_hash': 'نمک و هش scrypt رمز',
    'role': 'نقش admin یا manager یا staff', 'is_active': 'فعال بودن حساب یا خدمت',
    'token_hash': 'هش SHA256 توکن نشست', 'expires_at': 'زمان پایان اعتبار نشست',
    'user_id': 'شناسه کاربر', 'employee_id': 'شناسه کارمند', 'employee_name': 'نام کارمند در زمان ثبت',
    'action': 'نوع عملیات ثبت شده', 'entity': 'نوع موجودیت', 'entity_id': 'شناسه موجودیت تغییر یافته',
    'details': 'شرح عملیات بدون رمز و توکن', 'name': 'نام', 'description': 'شرح', 'notes': 'یادداشت',
    'email': 'ایمیل اختیاری', 'phone': 'تلفن اختیاری', 'address': 'نشانی', 'contact_name': 'نام رابط',
    'status': 'وضعیت تحت کنترل گردش کار', 'quantity': 'تعداد صحیح', 'amount': 'مبلغ صحیح به تومان',
    'unit_price': 'قیمت واحد به تومان', 'total': 'مجموع محاسبه شده به تومان', 'title': 'عنوان سند',
    'start_date': 'تاریخ شروع میلادی', 'end_date': 'تاریخ پایان میلادی', 'reason': 'دلیل ثبت',
    'reference': 'شماره مرجع یا پیگیری', 'type': 'نوع عملیات یا درخواست',
    'department_id': 'مرجع واحد سازمانی', 'customer_id': 'مرجع مشتری', 'contract_id': 'مرجع قرارداد',
    'item_id': 'مرجع کالا', 'asset_id': 'مرجع دارایی سریال دار', 'supplier_id': 'مرجع تأمین کننده',
    'quotation_id': 'مرجع پیشنهاد قیمت', 'invoice_id': 'مرجع صورتحساب', 'offering_id': 'مرجع خدمت کاتالوگ',
    'personnel_code': 'کد پرسنلی یکتا', 'first_name': 'نام کارمند', 'last_name': 'نام خانوادگی کارمند',
    'job_title': 'عنوان شغلی', 'employment_type': 'نوع همکاری', 'hire_date': 'تاریخ استخدام میلادی',
    'base_salary': 'حقوق پایه به تومان', 'allowance': 'فوق العاده مصوب مأموریت به تومان',
    'destination': 'مقصد مأموریت', 'kind': 'نوع تشویق یا جریمه', 'date': 'تاریخ تعلق تشویق یا جریمه',
    'period': 'ماه میلادی حقوق با قالب YYYY-MM', 'bonus': 'مزایای دستی فیش به تومان',
    'deduction': 'کسورات دستی فیش به تومان', 'comment': 'توضیح تصمیم مدیر',
    'sku': 'کد یکتای کالا', 'category': 'دسته بندی کالا', 'unit': 'واحد شمارش یا ارائه خدمت',
    'tracking_type': 'روش رهگیری تعدادی یا سریال دار', 'min_stock': 'آستانه هشدار موجودی',
    'serial_number': 'شماره سریال یکتای دستگاه', 'serial_numbers': 'آرایه شماره سریال های ورودی',
    'company_name': 'نام شرکت مشتری', 'valid_until': 'تاریخ پایان اعتبار پیشنهاد',
    'lines': 'آرایه ردیف های پیشنهاد طبق مدل ردیف ورودی همین راهنما', 'due_date': 'سررسید پرداخت صورتحساب',
    'issue_date': 'تاریخ صدور قطعی صورتحساب', 'paid_date': 'تاریخ دریافت ثبت شده',
    'method': 'روش ثبت دریافت وجه', 'password': 'گذرواژه اولیه یا ورود حساب',
    'current_password': 'گذرواژه فعلی برای تأیید تغییر', 'new_password': 'گذرواژه جدید حساب',
}
COMMON_ENDPOINTS = {
    'GET /health': ('سلامت اتصال دیتابیس و نام سامانه و حالت نمونه', 'بدون ورود', 'برای بررسی اجرا؛ شامل داده تجاری نیست'),
    'POST /api/v1/auth/login': ('ورود و دریافت توکن نشست', 'بدون ورود', '۱۰ تلاش در دقیقه برای هر IP در همین فرایند؛ نشست هشت ساعته'),
    'GET /api/v1/auth/me': ('مشخصات و نقش حساب فعلی', 'هر حساب فعال', 'توکن معتبر همان سامانه لازم است'),
    'POST /api/v1/auth/logout': ('خروج و ابطال توکن جاری', 'هر حساب فعال', 'توکن باطل شده دوباره قابل استفاده نیست'),
    'POST /api/v1/auth/change-password': ('تغییر رمز با ارائه رمز فعلی', 'هر حساب فعال', 'حداقل ۱۰ نویسه؛ همه نشست‌های کاربر باطل می‌شوند'),
    'GET /api/v1/users': ('فهرست کاربران و جست وجو', 'فقط admin', 'هش رمز در پاسخ حذف می‌شود'),
    'POST /api/v1/users': ('ایجاد حساب مستقل', 'فقط admin', 'نام کاربری یکتا؛ رمز حداقل ۱۰ نویسه'),
    'PATCH /api/v1/users/{record_id}': ('ویرایش نام و نقش و فعال بودن و مرجع کارمند', 'فقط admin', 'کاربر نمی‌تواند دسترسی مدیریت حساب جاری خودش را حذف کند'),
    'GET /api/v1/audit-logs': ('خواندن و جست وجوی سابقه عملیات', 'admin و manager', 'جست وجو و صفحه بندی؛ API حذف سابقه وجود ندارد'),
}


def xml_flag(parent, name, value='1'):
    element = OxmlElement(f'w:{name}')
    element.set(qn('w:val'), value)
    parent.append(element)
    return element


def direction(paragraph, rtl=True):
    ppr = paragraph._p.get_or_add_pPr()
    previous = ppr.find(qn('w:bidi'))
    if previous is not None:
        ppr.remove(previous)
    xml_flag(ppr, 'bidi', '1' if rtl else '0')
    # Word mirrors left/right paragraph alignment when bidi is enabled.
    # Logical start therefore uses LEFT for both RTL and LTR paragraphs.
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT


def format_run(run, latin=False, bold=None, size=None, monospace=None):
    family = 'Consolas' if (latin if monospace is None else monospace) else 'Tahoma'
    run.font.name = family
    rpr = run._r.get_or_add_rPr()
    rf = rpr.rFonts
    for key in ('ascii', 'hAnsi', 'cs', 'eastAsia'):
        rf.set(qn(f'w:{key}'), 'Tahoma' if key == 'cs' else family)
    lang = OxmlElement('w:lang')
    lang.set(qn('w:val'), 'en-US' if latin else 'fa-IR')
    lang.set(qn('w:bidi'), 'fa-IR')
    rpr.append(lang)
    xml_flag(rpr, 'rtl', '0' if latin else '1')
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold
        xml_flag(rpr, 'bCs', '1' if bold else '0')
    if size is not None:
        run.font.size = Pt(size)
        xml_flag(rpr, 'szCs', str(int(size * 2)))


def add_text(paragraph, text, size=None, bold=None, force_latin=False):
    # Code spans stay LTR inside the Persian paragraph.
    for index, fragment in enumerate(re.split(r'(`[^`]+`)', str(text))):
        if not fragment:
            continue
        latin = force_latin or (fragment.startswith('`') and fragment.endswith('`'))
        value = fragment[1:-1] if fragment.startswith('`') and fragment.endswith('`') else fragment
        # Give Persian inside JSON its own complex-script run, and keep raw
        # paths/identifiers LTR even when authors did not use code backticks.
        pattern = r'([\u0600-\u06ff\u200c]+(?:[ \u200c]+[\u0600-\u06ff\u200c]+)*)' if latin else r'''([A-Za-z0-9_./:{}?&=+<>@\\#%*"',\[\]()-]+(?:[ \t]+[A-Za-z0-9_./:{}?&=+<>@\\#%*"',\[\]()-]+)*)'''
        for token in re.split(pattern, value):
            if not token:
                continue
            is_farsi = bool(re.search(r'[\u0600-\u06ff]', token))
            is_latin = not is_farsi and (latin or bool(re.search(r'[A-Za-z0-9]', token)))
            run = paragraph.add_run(token)
            format_run(run, latin=is_latin, size=size, bold=bold, monospace=latin and is_latin)


def para(doc, text='', style=None, latin=False, size=None):
    p = doc.add_paragraph(style=style)
    direction(p, not latin)
    add_text(p, text, size=size, force_latin=latin)
    return p


def heading(doc, text, level=1):
    text = re.sub(r'[:؛|«»()\[\]؟!?،/\\–—-]', ' ', text)
    text = re.sub(r'\s+', ' ', text).strip()
    return para(doc, text, f'Heading {level}')


def code(doc, lines):
    if isinstance(lines, str):
        lines = lines.splitlines()
    for index, line in enumerate(lines):
        p = para(doc, line or ' ', 'Code', latin=True, size=9)
        p.paragraph_format.keep_with_next = index < len(lines) - 1


def cell_text(cell, text, header=False):
    p = cell.paragraphs[0]
    ascii_only = all(ord(c) < 128 for c in str(text))
    direction(p, not ascii_only)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.line_spacing = 1.1
    p.paragraph_format.keep_with_next = header
    add_text(p, str(text), size=9 if ascii_only and not header else 9.5, bold=header, force_latin=ascii_only)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def table(doc, headers, rows, widths=None):
    if not headers:
        return
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    total = 17.0
    widths = widths or [total / len(headers)] * len(headers)
    pr = t._tbl.tblPr
    xml_flag(pr, 'bidiVisual')
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        node = OxmlElement(f'w:{edge}')
        for k, v in {'val': 'single', 'sz': '5', 'color': 'D9D9D9'}.items():
            node.set(qn(f'w:{k}'), v)
        borders.append(node)
    pr.append(borders)
    margins = OxmlElement('w:tblCellMar')
    for edge in ('top', 'bottom', 'left', 'right'):
        n = OxmlElement(f'w:{edge}')
        n.set(qn('w:w'), '90')
        n.set(qn('w:type'), 'dxa')
        margins.append(n)
    pr.append(margins)
    for i, (c, title) in enumerate(zip(t.rows[0].cells, headers)):
        c.width = Cm(widths[i])
        cell_text(c, title, True)
        shading = OxmlElement('w:shd')
        shading.set(qn('w:fill'), 'E7EDF3')
        c._tc.get_or_add_tcPr().append(shading)
    xml_flag(t.rows[0]._tr.get_or_add_trPr(), 'tblHeader')
    for values in rows:
        row = t.add_row()
        xml_flag(row._tr.get_or_add_trPr(), 'cantSplit')
        for i, (c, value) in enumerate(zip(row.cells, values)):
            c.width = Cm(widths[i])
            cell_text(c, value)
    for gridcol, width in zip(t._tbl.tblGrid.gridCol_lst, widths):
        gridcol.set(qn('w:w'), str(int(Cm(width).twips)))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def picture(doc, path, caption, width=6.55):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))
    inline = run._r.xpath('.//wp:docPr')
    if inline:
        inline[0].set('descr', caption)
    c = para(doc, caption, 'Caption')
    c.paragraph_format.keep_with_next = False


def create_doc(service, kind):
    name, _, web, api = SERVICES[service]
    number, title = KINDS[kind]
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.7)
    sec.left_margin = sec.right_margin = Cm(1.9)
    sec.header_distance, sec.footer_distance = Cm(.7), Cm(.7)
    for style_name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Caption', 'Header', 'Footer'):
        style = doc.styles[style_name]
        style.font.name = 'Tahoma'
        style.font.color.rgb = RGBColor(0, 0, 0)
        rpr = style.element.get_or_add_rPr()
        fonts = OxmlElement('w:rFonts')
        for part in ('ascii', 'hAnsi', 'cs', 'eastAsia'):
            fonts.set(qn(f'w:{part}'), 'Tahoma')
        rpr.append(fonts)
        xml_flag(rpr, 'rtl')
    normal = doc.styles['Normal']
    normal.font.size = Pt(11)
    xml_flag(normal.element.get_or_add_rPr(), 'szCs', '22')
    normal.paragraph_format.line_spacing = 1.3
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.widow_control = True
    for style_name, size in [('Title', 22), ('Subtitle', 12), ('Heading 1', 15), ('Heading 2', 12), ('Caption', 9), ('Header', 8), ('Footer', 8)]:
        st = doc.styles[style_name]
        st.font.size = Pt(size)
        xml_flag(st.element.get_or_add_rPr(), 'szCs', str(size * 2))
        st.paragraph_format.space_before = Pt(14 if 'Heading' in style_name else 3)
        st.paragraph_format.space_after = Pt(7)
        if 'Heading' in style_name:
            st.paragraph_format.keep_with_next = True
    cs = doc.styles.add_style('Code', 1)
    cs.font.name = 'Consolas'
    cs.font.size = Pt(9)
    cs.paragraph_format.line_spacing = 1.05
    cs.paragraph_format.space_before = Pt(0)
    cs.paragraph_format.space_after = Pt(2)
    for s in (doc.styles['Title'], doc.styles['Subtitle']):
        ppr = s.element.get_or_add_pPr()
        for node in list(ppr):
            if node.tag in (qn('w:pBdr'), qn('w:shd')):
                ppr.remove(node)
    header = sec.header.paragraphs[0]
    direction(header)
    add_text(header, f'رویش   {name}   {title}', size=8)
    footer = sec.footer.paragraphs[0]
    direction(footer)
    add_text(footer, f'{name}   راهنمای {int(number)} از ۵   صفحه ', size=8)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    p = para(doc, f'{title} {name}', 'Title')
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    para(doc, 'شرکت توسعه نرم افزار رویش', 'Subtitle')
    para(doc, f'راهنمای {int(number)} از ۵   ویژه {name}   نسخه ۱   ۱ آبان ۱۴۰۵', size=9)
    doc.core_properties.title = f'{title} {name}'
    doc.core_properties.subject = f'راهنمای اختصاصی {name} در پروژه رویش'
    doc.core_properties.author = 'گروه پروژه رویش'
    doc.core_properties.keywords = f'رویش {name} {kind}'
    return doc


def font(size, bold=False):
    return ImageFont.truetype(str(FONT / ('arialbd.ttf' if bold else 'arial.ttf')), size)


def center(draw, box, text, size=28, bold=False, color='#111111'):
    x0, y0, x1, y1 = box
    f = font(size, bold)
    draw.text(((x0+x1)/2, (y0+y1)/2), text, font=f, fill=color, anchor='mm')


def arrow(draw, start, end, dashed=False, label=None):
    x0, y0 = start; x1, y1 = end
    if dashed:
        from math import hypot
        length = hypot(x1-x0, y1-y0)
        for at in range(0, int(length), 22):
            b = min(at+12, length)
            draw.line((x0+(x1-x0)*at/length,y0+(y1-y0)*at/length,x0+(x1-x0)*b/length,y0+(y1-y0)*b/length), fill='#687987', width=3)
    else:
        draw.line((start,end),fill='#516779',width=3)
    from math import atan2, cos, sin, pi
    angle = atan2(y1-y0, x1-x0)
    a = (x1-15*cos(angle-pi/7),y1-15*sin(angle-pi/7))
    b = (x1-15*cos(angle+pi/7),y1-15*sin(angle+pi/7))
    draw.polygon((end,a,b),fill='#516779')
    if label:
        px=(x0+x1)/2; py=(y0+y1)/2
        draw.rectangle((px-40,py-17,px+40,py+17),fill='white')
        draw.text((px,py),label,font=font(24),anchor='mm',fill='#111111')


def architecture_image(service, target):
    name, _, web, api = SERVICES[service]
    image = Image.new('RGB',(1500,1030),'white')
    d = ImageDraw.Draw(image)
    center(d,(0,10,1500,60),f'ROYESH  |  {service.upper()}  |  INDEPENDENT DEPLOYMENT',32,True)
    boxes = [((310,100,1190,235),'Browser runs React UI', 'Persian RTL UI / direct API calls with Bearer token'),
             ((55,385,720,555),f'Next.js   localhost:{web}',f'ROYESH_SYSTEM={service} / serves HTML CSS JS'),
             ((800,385,1450,555),f'FastAPI   localhost:{api}','/api/v1 / Auth / Validation / Business rules'),
             ((800,720,1450,880),f'{service}.db   SQLite','SQLAlchemy / PostgreSQL connection option')]
    for box,title,subtitle in boxes:
        d.rounded_rectangle(box,radius=8,fill='#F4F7F9',outline='#233F53',width=3)
        center(d,(box[0],box[1]+14,box[2],box[1]+67),title,32,True)
        center(d,(box[0],box[1]+73,box[2],box[3]-8),subtitle,24)
    for start,end,label in [((610,235),(390,385),'UI files'),((910,235),(1120,385),'HTTP JSON'),((1120,555),(1120,720),'SQL')]:
        arrow(d,start,end,label=label)
    d.line((160,935,1340,935), fill='#D9D9D9',width=2)
    center(d,(40,955,1460,1010),'Integration team: API gateway / SSO / cross-system queries (future)',26)
    image.save(target)


def box_anchor(box, other):
    x,y,w,h = box
    cx,cy=x+w/2,y+h/2
    ox,oy,ow,oh=other
    dx,dy=ox+ow/2-cx,oy+oh/2-cy
    ratio=min((w/2)/abs(dx) if dx else 1e9,(h/2)/abs(dy) if dy else 1e9)
    return cx+dx*ratio,cy+dy*ratio


def er_image(service, tables, target, common=False):
    selected = [t for t in tables if (t['name'] in COMMON_NOTES)==common]
    if common:
        positions={'users':(550,110,450,220),'auth_sessions':(40,475,450,220),'audit_logs':(1060,475,450,220)}
        height=840
    else:
        layouts={
            'personnel': {'hr_departments':(1,0),'hr_employees':(1,1),'hr_contracts':(0,1),'hr_leaves':(0,2),'hr_missions':(2,2),'hr_adjustments':(2,1),'hr_payrolls':(1,2)},
            'inventory': {'inv_suppliers':(0,0),'inv_items':(1,1),'inv_purchases':(2,0),'inv_assets':(0,2),'inv_assignments':(2,2),'inv_movements':(1,3)},
            'commerce': {'commerce_customers':(0,0),'commerce_offerings':(2,0),'commerce_quotations':(0,1),'commerce_quotation_lines':(2,1),'commerce_contracts':(0,2),'commerce_invoices':(1,3),'commerce_payments':(2,3)},
        }
        layout=layouts[service]
        positions={name:(45+x*525,100+y*305,455,220) for name,(x,y) in layout.items()}
        height=100+(max(y for x,y in layout.values())+1)*305+70
    image=Image.new('RGB',(1600,height),'white'); d=ImageDraw.Draw(image)
    center(d,(0,8,1600,64),f'{service.upper()}  |  '+('LOCAL AUTHENTICATION TABLES' if common else 'BUSINESS DATA RELATIONSHIPS'),32,True)
    for table_info in selected:
        for col in table_info['columns']:
            for ref in col['foreign_keys']:
                parent=ref.rsplit('.',1)[0]
                if parent in positions:
                    a,b=positions[parent],positions[table_info['name']]
                    child='0..1' if col['unique'] else 'N'
                    if service=='commerce' and parent=='commerce_customers' and table_info['name']=='commerce_contracts':
                        # Independent contracts also reference customers directly;
                        # route around the quotation box instead of hiding the FK.
                        d.line([(500,210),(590,210),(590,820)],fill='#516779',width=3)
                        arrow(d,(590,820),(500,820))
                        d.text((608,680),'1:N customer_id',font=font(24),fill='#111111')
                    elif service=='commerce' and parent=='commerce_invoices' and table_info['name']=='commerce_payments':
                        start,end=box_anchor(a,b),box_anchor(b,a)
                        arrow(d,start,end)
                        d.text(((start[0]+end[0])/2,start[1]-25),'1:N',font=font(24),anchor='mm',fill='#111111')
                    else:
                        arrow(d,box_anchor(a,b),box_anchor(b,a),label=f'1:{child}')
    if common:
        a,b=positions['users'],positions['audit_logs']
        arrow(d,box_anchor(a,b),box_anchor(b,a),dashed=True,label='1:N')
    for t in selected:
        x,y,w,h=positions[t['name']]
        d.rectangle((x,y,x+w,y+h),fill='white',outline='#294B63',width=3)
        d.rectangle((x+2,y+2,x+w-2,y+49),fill='#E7EDF3')
        center(d,(x,y,x+w,y+49),t['name'],26,True)
        priority=[c for c in t['columns'] if c['primary_key'] or c['foreign_keys']]
        more=[c for c in t['columns'] if c not in priority and c['name'] not in ('created_at','updated_at')]
        for i,c in enumerate((priority+more)[:5]):
            suffix=' PK' if c['primary_key'] else ' FK' if c['foreign_keys'] else ' UQ' if c['unique'] else ''
            d.text((x+15,y+63+i*29),c['name']+suffix,font=font(25),fill='#111111')
    center(d,(20,height-57,1580,height-10),'Solid = declared FK    Dashed = logical reference    PK = primary key    UQ = unique',24)
    image.save(target)


def authored(doc, sections):
    for section in sections:
        heading(doc, section['heading'])
        for text in section.get('paragraphs',[]):
            para(doc,text)
        for text in section.get('bullets',[]):
            para(doc,'• '+text)
        tb=section.get('table')
        if tb and tb.get('headers'):
            table(doc,tb['headers'],tb.get('rows',[]))
        if section.get('code'):
            code(doc,section['code'])


def db_dictionary(doc, data, contract):
    heading(doc,'فرهنگ جدول ها و ستون های واقعی')
    para(doc,'تمام جدول‌های زیر ستون‌های مشترک `id` و `created_at` و `updated_at` دارند. `id` رشته UUID و کلید اصلی است؛ دو زمان، DateTime با منطقه زمانی هستند و به UTC در API نمایش داده می‌شوند. این سه ستون برای کوتاه شدن جدول‌های زیر تکرار نشده‌اند. خالی بودن در این فرهنگ به معنی SQL NULL است؛ رشته خالی با NULL تفاوت دارد. پیش‌فرض‌های مدل در Python اجرا می‌شوند مگر قید server_default صریح وجود داشته باشد.')
    notes={**COMMON_NOTES,**data.get('table_notes',{})}
    for item in sorted(contract['tables'],key=lambda t:(t['name'] in COMMON_NOTES,t['name'])):
        heading(doc,item['name'],2)
        para(doc,notes.get(item['name'],'جدول اختصاصی همین سامانه'))
        rows=[]
        for col in item['columns']:
            if col['name'] in ('id','created_at','updated_at'):
                continue
            flags=[]
            if col['primary_key']: flags.append('PK')
            if col['unique']: flags.append('UNIQUE')
            flags.extend('FK '+ref for ref in col['foreign_keys'])
            if col['index']: flags.append('INDEX')
            note=data.get('field_notes',{}).get(item['name']+'.'+col['name'],FIELD_NOTES.get(col['name'],'فیلد اختصاصی جدول'))
            rows.append([col['name'],col['type'],'بله' if col['nullable'] else 'خیر','; '.join(flags) or '—',note])
        table(doc,['نام ستون','نوع ذخیره','NULL','کلید و نمایه','معنا'],rows,[4.0,2.6,1.1,3.5,5.8])
        for c in item['constraints']:
            if c['kind']=='UniqueConstraint' and ',' in c['value']:
                para(doc,'یکتایی ترکیبی ستون‌ها: `'+c['value']+'`',size=10)
            elif c['kind']=='CheckConstraint':
                para(doc,'قید بررسی دیتابیس',size=9)
                code(doc,[c['value']])
    heading(doc,'تطبیق با مهاجرت و نگهداری')
    para(doc,'مرجع کد مدل، فایل `models.py` همین سامانه و جدول‌های مشترک `shared/auth.py` و `shared/audit.py` است. مهاجرت `migrations/versions/001_initial.py` تصویر ثابت ساختار اولیه است. اصلاح مدل به تنهایی دیتابیس قدیمی را تغییر نمی‌دهد؛ برای تغییر ساختار باید revision جدید Alembic بسازید. اتصال مستقیم به جدول سامانه دیگر وجود ندارد و مرجع کارمند در سرویس‌های دیگر FK محلی نیست.')


def endpoint_catalog(doc,data,contract):
    heading(doc,'قرارداد عمومی API همین سامانه')
    _,_,web,api=SERVICES[data['service']]
    code(doc,[f'Base URL: http://localhost:{api}/api/v1',f'Swagger: http://localhost:{api}/docs',f'OpenAPI: http://localhost:{api}/openapi.json','Authorization: Bearer <access_token>','Content-Type: application/json'])
    para(doc,'مبالغ عدد صحیح تومان، شناسه‌ها UUID و تاریخ ورودی میلادی `YYYY-MM-DD` است. زمان‌های ثبت ISO 8601 به UTC هستند. فهرست‌های صفحه‌بندی‌شده، `items` و `total` و `page` و `page_size` دارند. صفحه از ۱ و اندازه صفحه از ۱ تا ۱۰۰ است. فیلتر هر مسیر در فهرست زیر آمده است؛ نباید فرض کرد تمام مسیرها فیلتر یکسان دارند. کلید ناشناخته ورودی رد می‌شود. در PATCH تنها فیلدهای تغییرکرده را بفرستید؛ NULL فقط برای فیلدهای مجاز پذیرفته می‌شود.')
    table(doc,['کد','معنی و اقدام'],[
        ['200 / 201','عملیات موفق؛ پاسخ واقعی API مبنای نمایش رابط است'],['400','درخواست نامعتبر مانند رمز فعلی نادرست'],['401','ورود لازم یا نشست نامعتبر یا منقضی؛ دوباره وارد شوید'],['403','نقش حساب برای این عملیات مجاز نیست'],['404','رکورد پیدا نشد؛ شناسه و سامانه را کنترل کنید'],['409','تعارض وضعیت یا یکتایی یا مانده؛ ابتدا آخرین سابقه را بخوانید'],['422','نوع یا مقدار ورودی یا قاعده تاریخ معتبر نیست'],['429','تلاش ورود زیاد؛ یک دقیقه صبر کنید']], [2.5,14.5])
    notes=data.get('endpoint_notes',{})
    groups={}
    for path,operations in contract['openapi']['paths'].items():
        for method,op in operations.items():
            if method not in ('get','post','patch','delete','put'):
                continue
            key=method.upper()+' '+path
            common=COMMON_ENDPOINTS.get(key)
            info=notes.get(key,{})
            if not info and common:
                info={'purpose':common[0],'access':common[1],'rules':common[2]}
            group='حساب و سلامت و سابقه' if common else 'مسیرهای اختصاصی کسب و کار'
            groups.setdefault(group,[]).append((path,method,op,info))
    for group,rows in groups.items():
        heading(doc,group)
        for path,method,op,info in rows:
            p=para(doc,method.upper()+' '+path,'Code',latin=True,size=9.5)
            p.paragraph_format.keep_with_next=True
            purpose=info.get('purpose') or op.get('summary','')
            para(doc,purpose+'؛ دسترسی: '+info.get('access','مطابق نقش و مالکیت رکورد')+'؛ '+info.get('rules',''),size=10)
            params=[]
            for param in op.get('parameters',[]):
                schema=param.get('schema',{})
                typ=schema_type(schema)
                extra=('الزامی' if param.get('required') else 'اختیاری')
                params.append('`'+param['name']+'` از نوع `'+typ+'` '+extra)
            if params:
                para(doc,'پارامترها: '+'؛ '.join(params),size=9)
            schema=op.get('requestBody',{}).get('content',{}).get('application/json',{}).get('schema',{})
            if schema:
                para(doc,'مدل بدنه: `'+schema_type(schema)+'`؛ فیلدهای آن در فرهنگ ورودی انتهای راهنما آمده است.',size=9)
    heading(doc,'نمونه ورود و استفاده از نشست')
    code(doc,['POST /api/v1/auth/login',json.dumps({'username':'admin','password':'admin12345'},indent=2,ensure_ascii=False),'', 'GET /api/v1/auth/me', 'Authorization: Bearer <access_token>'])
    para(doc,'حساب نمونه بالا فقط برای محیط ساختگی `SEED_DEMO=true` است. مقدار `access_token` را از پاسخ ورود بگیرید و در سربرگ درخواست‌های بعدی همان سامانه قرار دهید. توکن سرویس دیگر معتبر نیست. رمز و توکن را در URL، فایل عمومی یا سابقه عملیات ثبت نکنید. خروج و تغییر رمز، اعتبار نشست را پایان می‌دهد.')
    for example in data.get('examples',[]):
        heading(doc,example['title'],2)
        para(doc,example.get('explanation','شناسه‌ها را از پاسخ عملیات قبلی جایگزین کنید.'))
        lines=[example['method']+' '+example['path']]
        if example.get('body') is not None:
            lines.extend(json.dumps(example['body'],ensure_ascii=False,indent=2).splitlines())
        code(doc,lines)
    heading(doc,'فرهنگ مدل های ورودی')
    para(doc,'این جدول‌ها از قرارداد OpenAPI برنامه تولید شده‌اند. علامت الزامی برای مدل همان درخواست است؛ اختیاری بودن فیلد در مدل PATCH به معنی پذیرفته شدن هر مقدار یا امکان تغییر سند نهایی نیست. محدودیت وضعیت و مانده در منطق بک‌اند نیز کنترل می‌شود. مدل‌های پاسخ عمومی این نسخه همیشه در OpenAPI به صورت تفصیلی تایپ نشده‌اند؛ برای شکل پاسخ، خروجی واقعی Swagger و توضیح گردش را مبنا قرار دهید.')
    schemas=contract['openapi'].get('components',{}).get('schemas',{})
    for name,schema in schemas.items():
        if name in ('ValidationError','HTTPValidationError') or not schema.get('properties'):
            continue
        heading(doc,name,2)
        required=set(schema.get('required',[]))
        values=[]
        for key,field in schema['properties'].items():
            details=[]
            for k,label in [('minimum','حداقل'),('maximum','حداکثر'),('exclusiveMinimum','بیشتر از'),('minLength','حداقل طول'),('maxLength','حداکثر طول'),('minItems','حداقل ردیف'),('maxItems','حداکثر ردیف')]:
                if k in field: details.append(f'{label} {field[k]}')
            if 'default' in field:
                details.append('پیش فرض '+str(field['default']))
            if 'pattern' in field:
                details.append('الگوی قالب در Swagger')
            meaning=FIELD_NOTES.get(key,'فیلد اختصاصی این مدل')
            values.append([key,schema_type(field),'بله' if key in required else 'خیر',meaning+'؛ '+'، '.join(details) if details else meaning])
        table(doc,['فیلد','نوع و مقدار مجاز','الزامی','شرح و محدودیت'],values,[3.3,4.4,1.5,7.8])


def schema_type(schema):
    if '$ref' in schema:
        return schema['$ref'].rsplit('/',1)[-1]
    if 'anyOf' in schema or 'oneOf' in schema:
        return ' | '.join(schema_type(x) for x in schema.get('anyOf',schema.get('oneOf',[])))
    if 'enum' in schema:
        return ' / '.join(str(x) for x in schema['enum'])
    typ=schema.get('type','object')
    if typ=='array':
        return 'array['+schema_type(schema.get('items',{}))+']'
    return schema.get('format',typ)


def frontend_catalog(doc,service):
    config=json.loads((QA/'contracts/frontend.json').read_text(encoding='utf-8'))
    heading(doc,'نقشه صفحه ها و اتصال به API')
    para(doc,'این رابط یک صفحه Next.js و چند بخش با hash دارد. برای نمونه `/#dashboard` نمای کلی را باز می‌کند؛ بخش‌ها فایل route جدا در App Router ندارند. جدول زیر فقط تنظیمات همین سامانه را از `resources.ts` نشان می‌دهد. مشخصات منو و محدودیت نقش در `App.tsx` و کنترل واقعی مجوز در API اعمال می‌شود.')
    rows=[]
    for r in config['resources'][service]:
        gate='فقط مدیر و کارشناس' if r.get('managerOnly') else 'مطابق نقش و مالکیت'
        if r.get('readOnly'): gate+='؛ فقط خواندنی'
        rows.append([r['title'],'/#'+r['key'],r['endpoint'],gate])
    table(doc,['صفحه','نشانی بخش','API پایه','دسترسی'],rows,[3.5,3.8,4.2,5.5])
    heading(doc,'تعریف فیلدهای هر بخش')
    for r in config['resources'][service]:
        heading(doc,r['title'],2)
        para(doc,r['description'])
        if r.get('fields'):
            rows=[]
            for f in r['fields']:
                rows.append([f['key'],f['label'],f.get('type','text'),('فقط ایجاد' if f.get('createOnly') else 'مطابق گردش')+('؛ مرجع '+f['reference'] if f.get('reference') else '')])
            table(doc,['کلید','برچسب رابط','کنترل','رفتار'],rows,[3.3,4.3,2.5,6.9])
        para(doc,'ستون‌های فهرست: '+'، '.join(c['label'] for c in r['columns'])+'؛ وضعیت‌ها و اقدام‌های نهایی، از تعریف همین بخش و `workflows.ts` گرفته می‌شوند.',size=10)


def source_end(doc,service,kind):
    heading(doc,'مرجع مطالعه و کنترل نسخه')
    paths={'architecture':['README.md','shared/app.py','scripts/manage.py','compose.yaml','docs/integration.fa.md'],
           'database':[f'services/{service}/models.py','shared/auth.py','shared/audit.py','migrations/versions/001_initial.py'],
           'backend':[f'services/{service}/app.py',f'services/{service}/schemas.py','shared/database.py','shared/auth.py','tests'],
           'frontend':['frontend/src/app/page.tsx','frontend/src/components/App.tsx','frontend/src/lib/resources.ts','frontend/src/lib/workflows.ts','frontend/tests/frontend.test.mjs'],
           'user_api':['README.md',f'docs/openapi/{service}.json','docs/api-contract.md','docs/verification.fa.md']}[kind]
    para(doc,'مسیرهای زیر نسبت به ریشه پروژه هستند. برای بررسی رفتار تازه، کد و OpenAPI همان نسخه اجراشده را بخوانید. پس از تغییر API، فایل قرارداد و این راهنما نیز باید به روز شود. این مجموعه به همین سامانه اختصاص دارد؛ حساب‌ها و پایگاه داده دو سامانه دیگر مستقل‌اند.')
    code(doc,paths)


def main():
    import argparse
    parser=argparse.ArgumentParser()
    parser.add_argument('--service',choices=list(SERVICES))
    parser.add_argument('--kind',choices=list(KINDS))
    args=parser.parse_args()
    manifest=[]
    for service,(name,folder,web,api) in SERVICES.items():
        if args.service and args.service!=service: continue
        data=json.loads((SOURCE/f'{service}.json').read_text(encoding='utf-8-sig'))
        assert data['service']==service
        contract=json.loads((QA/'contracts'/f'{service}.json').read_text(encoding='utf-8'))
        assets=QA/'diagrams'/service
        assets.mkdir(parents=True,exist_ok=True)
        architecture_image(service,assets/'architecture.png')
        er_image(service,contract['tables'],assets/'database.png')
        er_image(service,contract['tables'],assets/'auth.png',common=True)
        target=OUTPUT/folder
        target.mkdir(parents=True,exist_ok=True)
        for kind,(number,title) in KINDS.items():
            if args.kind and args.kind!=kind: continue
            content=data['documents'][kind]
            doc=create_doc(service,kind)
            para(doc,content['intro'])
            if kind=='backend':
                para(doc,'فرمان اجرای مستقل از ریشه پروژه')
                code(doc,[f'python -m uvicorn services.{service}.app:app --port {api}'])
            if kind=='architecture':
                picture(doc,assets/'architecture.png',f'شکل ۱ معماری اجرا و مالکیت داده در {name}')
                para(doc,'پیکان‌ها مسیر درخواست را نشان می‌دهند. HTTP JSON بین مرورگر و API و SQL درون مرز همین سامانه است. سطر پایین نمودار، مسئولیت آتی تیم ادغام را معرفی می‌کند و جزء اتصال پیاده‌سازی‌شده فعلی نیست.',size=10)
            if kind=='database':
                picture(doc,assets/'database.png',f'شکل ۱ روابط جدول های کسب و کار {name}')
                para(doc,'در نمودار، PK کلید اصلی و FK کلید خارجی واقعی است. خط پیوسته FK اعلام‌شده را نشان می‌دهد؛ جهت پیکان از جدول مرجع به جدول دارای FK است. N یعنی چند رکورد وابسته و 0..1 یعنی رابطه یکتای اختیاری. NULL و قیود دقیق هر ستون در فرهنگ پایین آمده‌اند. نمودار برای خوانایی تنها کلیدها و چند ستون مهم را نشان می‌دهد.',size=10)
            authored(doc,content['sections'])
            if kind=='database':
                heading(doc,'جدول های حساب و سابقه در همین دیتابیس')
                picture(doc,assets/'auth.png',f'شکل ۲ جدول های احراز هویت مستقل {name}')
                para(doc,'ارتباط نشست با کاربر، FK واقعی است. ارتباط رویداد با کاربر منطقی و متنی است و با خط چین نشان داده می‌شود؛ این تفاوت برای حفظ سابقه مهم است. مرجع کارمند حساب نیز FK ندارد؛ در سامانه پرسنلی وجود پرونده در API بررسی می‌شود.')
                db_dictionary(doc,data,contract)
            if kind=='frontend': frontend_catalog(doc,service)
            if kind=='user_api': endpoint_catalog(doc,data,contract)
            source_end(doc,service,kind)
            path=target/f'{service}-{number}-{kind.replace("_","-")}.docx'
            doc.save(path)
            manifest.append({'service':service,'kind':kind,'title':f'{title} {name}','path':str(path.relative_to(ROOT))})
            print(path,flush=True)
    QA.mkdir(parents=True,exist_ok=True)
    (QA/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')


if __name__=='__main__':
    main()
